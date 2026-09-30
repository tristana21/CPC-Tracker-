#!/usr/bin/env python3
"""
Name Flagger: a web app that crawls websites (and the links on them), or scans
documents you upload (PDF, Word, text), and flags every place a name from your
list is mentioned.

Run:
    pip install -r requirements.txt
    python app.py
Then open http://127.0.0.1:5000 in your browser.
"""
import csv
import ipaddress
import os
import socket
import io
import re
import threading
import time
import uuid
from collections import deque
from urllib import robotparser
from urllib.parse import unquote, urldefrag, urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from docx import Document
from flask import Flask, Response, jsonify, render_template_string, request
from pypdf import PdfReader

app = Flask(__name__)

USER_AGENT = "NameFlagger/1.0 (personal research crawler; respects robots.txt)"
MAX_BYTES = 5_000_000          # skip pages larger than ~5 MB
SNIPPETS_PER_HIT = 3           # context snippets kept per (name, page)
CONTEXT_CHARS = 90             # characters of context on each side of a match
SKIP_EXT = (".pdf", ".jpg", ".jpeg", ".png", ".gif", ".webp", ".svg", ".zip", ".gz",
            ".mp4", ".mp3", ".mov", ".avi", ".doc", ".docx", ".xls", ".xlsx",
            ".ppt", ".pptx", ".exe", ".dmg", ".css", ".js", ".ico", ".woff", ".woff2")

JOBS = {}
ALLOW_PRIVATE = os.environ.get("ALLOW_PRIVATE") == "1"   # set to 1 only for local testing
MAX_RUNNING = int(os.environ.get("MAX_RUNNING", "3"))
PAGE_CAP = int(os.environ.get("PAGE_CAP", "1000"))
MAX_FILES = int(os.environ.get("MAX_FILES", "20"))
MAX_UPLOAD_MB = int(os.environ.get("MAX_UPLOAD_MB", "25"))
DOC_EXT = (".pdf", ".docx", ".txt", ".md", ".csv", ".html", ".htm")
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_MB * 1024 * 1024


def is_public(url):
    """Block crawling of localhost / private network addresses (important once deployed)."""
    if ALLOW_PRIVATE:
        return True
    host = urlparse(url).hostname
    if not host:
        return False
    try:
        infos = socket.getaddrinfo(host, None)
    except socket.gaierror:
        return False
    return all(ipaddress.ip_address(i[4][0]).is_global for i in infos)

# ----------------------------------------------------------------- matching --
B_START, B_END = r"(?<!\w)", r"(?!\w)"
SP = r"\s+"


def _tok(t):
    """Escape a name part, tolerating hyphen/space and straight/curly apostrophes."""
    return re.escape(t).replace(r"\-", r"[-\s]?").replace("'", "['’]?")


def build_variants(alias):
    """Return (label, regex) pairs for common ways a name is written."""
    alias = alias.replace("’", "'").strip()
    parts = alias.split()
    if len(parts) == 1:
        return [("exact", B_START + _tok(parts[0]) + B_END)]

    first, last = parts[0], parts[-1]
    F, L = _tok(first), _tok(last)
    initial = r"(?-i:" + re.escape(first[0].upper()) + r")\.?"
    middle = r"(?:(?-i:[A-Z])\.?|(?-i:[A-Z])[\w'’-]+)"   # capitalised middle name/initial

    v = [("full name", SP.join(_tok(p) for p in parts))]
    if len(parts) > 2:
        v.append(("first + last", F + SP + L))
    v += [
        ("with middle name/initial", F + SP + middle + SP + L),
        ("Last, First", L + r",\s*" + F + r"(?:\s+(?-i:[A-Z])\.?)?"),
        ("initial + last", initial + SP + L),
    ]
    return [(label, B_START + pat + B_END) for label, pat in v]


def parse_names(raw):
    """One name per line; aliases on the same line separated by |."""
    names = []
    for line in raw.splitlines():
        aliases = [a.strip() for a in line.split("|") if a.strip()]
        if not aliases:
            continue
        seen, variants = set(), []
        for a in aliases:
            for label, pat in build_variants(a):
                if pat in seen:
                    continue
                seen.add(pat)
                lab = label if a == aliases[0] else f"alias: {a} ({label})"
                variants.append((lab, re.compile(pat, re.IGNORECASE)))
        names.append({"name": aliases[0], "variants": variants})
    return names


def find_mentions(text, variants):
    """All non-overlapping matches across variants, preferring the longest."""
    spans = []
    for label, rx in variants:
        spans += [(m.start(), m.end(), label) for m in rx.finditer(text)]
    spans.sort(key=lambda s: (s[0], -(s[1] - s[0])))
    out, last_end = [], -1
    for s in spans:
        if s[0] >= last_end:
            out.append(s)
            last_end = s[1]
    return out


def url_to_words(s):
    return re.sub(r"[-_./+=?&%#:]+", " ", unquote(s))


# ------------------------------------------------------------------ crawling --
def robots_allowed(session, cache, url):
    p = urlparse(url)
    base = f"{p.scheme}://{p.netloc}"
    if base not in cache:
        rp = robotparser.RobotFileParser()
        try:
            r = session.get(base + "/robots.txt", timeout=8)
            if r.status_code in (401, 403):
                rp.disallow_all = True
            elif r.status_code >= 400:
                rp.allow_all = True
            else:
                rp.parse(r.text.splitlines())
        except requests.RequestException:
            rp.allow_all = True
        cache[base] = rp
    return cache[base].can_fetch(USER_AGENT, url)


def record(job, name, page_url, title, label, where, snippet, is_file=False):
    key = (name, page_url)
    h = job["hits"].get(key)
    if h is None:
        h = {"name": name, "url": page_url, "title": title, "count": 0, "file": is_file,
             "variants": [], "where": [], "snippets": []}
        job["hits"][key] = h
        job["order"].append(h)
    h["count"] += 1
    if label not in h["variants"]:
        h["variants"].append(label)
    if where not in h["where"]:
        h["where"].append(where)
    if len(h["snippets"]) < SNIPPETS_PER_HIT:
        h["snippets"].append(snippet)


def text_snippet(text, start, end):
    a, b = max(0, start - CONTEXT_CHARS), min(len(text), end + CONTEXT_CHARS)
    return {"pre": ("…" if a else "") + text[a:start], "match": text[start:end],
            "post": text[end:b] + ("…" if b < len(text) else ""), "where": "page text"}


# ------------------------------------------------------------------ documents --
def extract_sections(filename, data):
    """Return [(where, text)] for an uploaded document. PDFs are split by page."""
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("PDF is password-protected")
        return [(f"page {i}", p.extract_text() or "") for i, p in enumerate(reader.pages, 1)]
    if ext == ".docx":
        d = Document(io.BytesIO(data))
        parts = [para.text for para in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" ".join(c.text for c in row.cells))
        for sec in d.sections:
            for part in (sec.header, sec.footer):
                parts += [para.text for para in part.paragraphs]
        return [("document", "\n".join(parts))]
    if ext == ".doc":
        raise ValueError("Old .doc files aren't supported. Save it as .docx or PDF and try again")
    if ext not in DOC_EXT:
        raise ValueError("Unsupported file type. Use PDF, Word (.docx), or text files")
    text = data.decode("utf-8", errors="replace")
    if ext in (".html", ".htm"):
        text = BeautifulSoup(text, "html.parser").get_text(" ")
    return [("document", text)]


def scan_files(job):
    names = job["names"]
    while job["files"] and not job["stop"]:
        filename, data = job["files"].pop(0)
        job["current"] = filename
        try:
            sections = extract_sections(filename, data)
        except Exception as e:
            msg = str(e) if isinstance(e, ValueError) else f"Couldn't read this file ({type(e).__name__})"
            job["errors"].append({"url": filename, "error": msg[:200]})
            job["files_scanned"] += 1
            continue
        before = len(job["order"])
        found_text = False
        for where, raw in sections:
            text = " ".join(raw.split())
            if not text:
                continue
            found_text = True
            for n in names:
                for st, en, label in find_mentions(text, n["variants"]):
                    snip = text_snippet(text, st, en)
                    snip["where"] = where
                    record(job, n["name"], filename, filename, label, where, snip, is_file=True)
        if not found_text:
            job["errors"].append({"url": filename,
                                  "error": "No readable text found (it may be a scanned image)"})
        job["files_scanned"] += 1
        job["recent"].appendleft({"url": filename, "title": filename, "file": True,
                                  "new_flags": len(job["order"]) - before})


def run_job(job):
    try:
        scan_files(job)
    except Exception as e:
        job["errors"].append({"url": job.get("current", ""), "error": repr(e)[:300]})
    job["files"] = []
    crawl(job)


def crawl(job):
    cfg, names = job["cfg"], job["names"]
    s = requests.Session()
    s.headers.update({"User-Agent": USER_AGENT, "Accept": "text/html,application/xhtml+xml"})
    robots_cache, last_request = {}, {}
    start_hosts = {urlparse(u).netloc.lower() for u in cfg["start_urls"]}
    queue = deque((u, 0) for u in cfg["start_urls"])
    seen = set(cfg["start_urls"])
    seen_cap = cfg["max_pages"] * 25

    try:
        while queue and not job["stop"] and job["pages_crawled"] < cfg["max_pages"]:
            url, depth = queue.popleft()
            job["queued"], job["current"] = len(queue), url

            if not is_public(url):
                job["errors"].append({"url": url, "error": "Skipped: private or unreachable address"})
                continue
            if cfg["respect_robots"] and not robots_allowed(s, robots_cache, url):
                job["skipped_robots"] += 1
                continue

            host = urlparse(url).netloc.lower()          # polite per-host delay
            wait = cfg["delay"] - (time.time() - last_request.get(host, 0))
            if wait > 0:
                time.sleep(wait)
            last_request[host] = time.time()

            try:
                r = s.get(url, timeout=12, stream=True)
                ctype = r.headers.get("Content-Type", "").lower()
                too_big = int(r.headers.get("Content-Length") or 0) > MAX_BYTES
                if "html" not in ctype or too_big:
                    r.close()
                    continue
                if r.status_code >= 400:
                    job["errors"].append({"url": url, "error": f"HTTP {r.status_code}"})
                    r.close()
                    continue
                html = r.text
            except requests.RequestException as e:
                job["errors"].append({"url": url, "error": str(e)[:200]})
                continue

            page = urldefrag(r.url)[0]
            if not is_public(page):
                continue
            job["pages_crawled"] += 1
            before = len(job["order"])
            soup = BeautifulSoup(html, "html.parser")
            title = soup.title.get_text(" ", strip=True)[:200] if soup.title else ""
            links = [(a["href"], " ".join(a.get_text(" ").split()))
                     for a in soup.find_all("a", href=True)]
            for t in soup(["script", "style", "noscript", "template", "svg", "title", "head"]):
                t.decompose()
            text = " ".join(soup.get_text(" ").split())

            # 1) names in the visible page text
            for n in names:
                for st, en, label in find_mentions(text, n["variants"]):
                    record(job, n["name"], page, title, label, "page text",
                           text_snippet(text, st, en))

            # 2) names hidden in link URLs (e.g. /people/jane-doe) + queue links
            for href, anchor in links:
                link = urldefrag(urljoin(page, href.strip()))[0]
                p = urlparse(link)
                if p.scheme not in ("http", "https"):
                    continue
                words = url_to_words(p.path + " " + p.query)
                for n in names:
                    if find_mentions(anchor, n["variants"]):
                        continue  # already counted via page text
                    if find_mentions(words, n["variants"]):
                        record(job, n["name"], page, title, "name in link URL", "link URL",
                               {"pre": "Links to ", "match": link,
                                "post": f' (link text: "{anchor[:80]}")' if anchor else "",
                                "where": "link URL"})

                if depth >= cfg["max_depth"] or link in seen or len(seen) >= seen_cap:
                    continue
                if not cfg["follow_external"] and p.netloc.lower() not in start_hosts:
                    continue
                if p.path.lower().endswith(SKIP_EXT):
                    continue
                seen.add(link)
                queue.append((link, depth + 1))

            job["recent"].appendleft({"url": page, "title": title,
                                      "new_flags": len(job["order"]) - before})

        job["status"] = "stopped" if job["stop"] else "done"
    except Exception as e:  # keep the UI informed instead of dying silently
        job["status"] = "error"
        job["errors"].append({"url": job.get("current", ""), "error": repr(e)[:300]})
    finally:
        job["current"], job["finished"] = "", time.time()


# -------------------------------------------------------------------- routes --
def clamp(v, lo, hi, default):
    try:
        return max(lo, min(hi, type(default)(v)))
    except (TypeError, ValueError):
        return default


@app.get("/")
def index():
    return render_template_string(PAGE, max_files=MAX_FILES, max_mb=MAX_UPLOAD_MB)


@app.errorhandler(413)
def too_large(_):
    return jsonify(error=f"Those files are too big together. The limit is {MAX_UPLOAD_MB} MB per scan."), 413


@app.post("/start")
def start():
    if request.content_type and request.content_type.startswith("multipart/"):
        d = request.form
        uploads = [f for f in request.files.getlist("files") if f and f.filename]
    else:
        d = request.get_json(force=True) or {}
        uploads = []
    if len(uploads) > MAX_FILES:
        return jsonify(error=f"You can scan up to {MAX_FILES} files at a time."), 400
    files, used = [], set()
    for f in uploads:
        fname = os.path.basename(f.filename.replace("\\", "/"))[:150] or "file"
        base, k = fname, 2
        while fname in used:
            stem, ext = os.path.splitext(base)
            fname, k = f"{stem} ({k}){ext}", k + 1
        used.add(fname)
        files.append((fname, f.read()))
    urls = []
    for u in d.get("urls", "").split():
        if not re.match(r"https?://", u, re.I):
            u = "https://" + u
        u = urldefrag(u)[0]
        if u not in urls:
            urls.append(u)
    names = parse_names(d.get("names", ""))
    if not (urls or files) or not names:
        return jsonify(error="Add a starting URL or a document, plus at least one name."), 400
    as_bool = lambda v, default: default if v is None else (v if isinstance(v, bool) else str(v).lower() in ("1", "true", "on", "yes"))

    cfg = {
        "start_urls": urls,
        "max_depth": clamp(d.get("max_depth"), 0, 6, 2),
        "max_pages": clamp(d.get("max_pages"), 1, PAGE_CAP, 200),
        "delay": clamp(d.get("delay"), 0.0, 10.0, 0.5),
        "follow_external": as_bool(d.get("follow_external"), True),
        "respect_robots": as_bool(d.get("respect_robots"), True),
    }
    if sum(j["status"] == "running" for j in JOBS.values()) >= MAX_RUNNING:
        return jsonify(error="The server is busy with other crawls. Try again in a minute."), 429
    job = {"id": uuid.uuid4().hex[:10], "recent": deque(maxlen=8), "status": "running", "cfg": cfg, "names": names,
           "pages_crawled": 0, "queued": len(urls), "current": "", "errors": [],
           "skipped_robots": 0, "hits": {}, "order": [], "stop": False,
           "files": files, "files_total": len(files), "files_scanned": 0,
           "started": time.time(), "finished": None}
    JOBS[job["id"]] = job
    threading.Thread(target=run_job, args=(job,), daemon=True).start()
    return jsonify(job_id=job["id"])


@app.get("/status/<job_id>")
def status(job_id):
    job = JOBS.get(job_id)
    if not job:
        return jsonify(error="No crawl with that ID. Start a new one."), 404
    hits = [dict(h, variants=list(h["variants"]), where=list(h["where"]),
                 snippets=list(h["snippets"])) for h in list(job["order"])]
    return jsonify(
        status=job["status"], pages_crawled=job["pages_crawled"], queued=job["queued"],
        max_pages=job["cfg"]["max_pages"], current=job["current"],
        files_total=job["files_total"], files_scanned=job["files_scanned"],
        progress_total=job["files_total"] + (job["cfg"]["max_pages"] if job["cfg"]["start_urls"] else 0),
        skipped_robots=job["skipped_robots"], errors=job["errors"][-50:],
        error_count=len(job["errors"]), names=[n["name"] for n in job["names"]],
        recent=list(job["recent"]), start_urls=job["cfg"]["start_urls"],
        elapsed=round((job["finished"] or time.time()) - job["started"], 1), hits=hits)


@app.post("/stop/<job_id>")
def stop(job_id):
    if job_id in JOBS:
        JOBS[job_id]["stop"] = True
    return jsonify(ok=True)


@app.get("/export/<job_id>.csv")
def export(job_id):
    job = JOBS.get(job_id)
    if not job:
        return "Unknown crawl", 404
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["name", "page_title", "url_or_file", "mentions", "matched_as", "found_in", "context"])
    for h in list(job["order"]):
        ctx = " | ".join(s["pre"] + "[" + s["match"] + "]" + s["post"] for s in h["snippets"])
        w.writerow([h["name"], h["title"], h["url"], h["count"],
                    "; ".join(h["variants"]), "; ".join(h["where"]), ctx])
    return Response(buf.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": f"attachment; filename=name-flags-{job_id}.csv"})


# ---------------------------------------------------------------------- page --
PAGE = r"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Name Flagger</title>
<link rel="icon" href="data:image/svg+xml,<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 100 100'><text y='.9em' font-size='90'>🚩</text></svg>">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Public+Sans:wght@400;500;600;750&display=swap" rel="stylesheet">
<style>
:root{
  --bg:#F2F3F6; --surface:#FFFFFF; --sunk:#E9EBF0; --ink:#151B28; --muted:#5E6778; --line:#DCE0E8;
  --accent:#3654E0; --accent-ink:#FFFFFF; --accent-soft:#E5EAFD;
  --mark:#FFE45C; --mark-ink:#151B28; --warn:#B4412E; --ok:#1F7A4D;
  --shadow:0 1px 2px rgba(21,27,40,.06),0 8px 24px -12px rgba(21,27,40,.18);
  font-family:"Public Sans",system-ui,-apple-system,"Segoe UI",Roboto,sans-serif;
  color:var(--ink); background:var(--bg); color-scheme:light;
}
@media (prefers-color-scheme:dark){:root{
  --bg:#10151F; --surface:#18202D; --sunk:#131A25; --ink:#E7EBF2; --muted:#98A2B4; --line:#283244;
  --accent:#86A2FF; --accent-ink:#0D1320; --accent-soft:#1E2A47;
  --mark:#E9CB30; --mark-ink:#10151F; --warn:#F2907A; --ok:#6FD39E;
  --shadow:0 1px 2px rgba(0,0,0,.3),0 10px 30px -14px rgba(0,0,0,.6); color-scheme:dark;
}}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);line-height:1.5;font-size:15px}
button,input,select,textarea{font:inherit;color:inherit}
:focus-visible{outline:2px solid var(--accent);outline-offset:2px}
.shell{max-width:1180px;margin:0 auto;padding:1.5rem 1.25rem 4rem}

/* header */
header{display:flex;align-items:center;justify-content:space-between;gap:1rem;margin-bottom:1.75rem}
.brand{display:flex;align-items:center;gap:.7rem}
.logo{width:38px;height:38px;border-radius:10px;background:#151B28;border:1px solid var(--line);display:grid;place-items:center}
.logo svg{width:20px;height:20px}
.brand h1{font-size:1.2rem;margin:0;font-weight:750;letter-spacing:-.01em}
.brand p{margin:0;color:var(--muted);font-size:.85rem}
.pill{display:inline-flex;align-items:center;gap:.45rem;font-size:.82rem;font-weight:600;padding:.35rem .75rem;border-radius:999px;background:var(--sunk);color:var(--muted)}
.pill .dot{width:8px;height:8px;border-radius:50%;background:currentColor}
.pill.running{background:var(--accent-soft);color:var(--accent)}
.pill.running .dot{animation:pulse 1.2s ease-in-out infinite}
.pill.done{color:var(--ok)}
.pill.error{color:var(--warn)}
@keyframes pulse{50%{opacity:.25}}

/* setup */
.setup{background:var(--surface);border:1px solid var(--line);border-radius:14px;box-shadow:var(--shadow);padding:1.4rem}
.headline{font-size:clamp(1.5rem,3vw,2.1rem);line-height:1.15;font-weight:750;letter-spacing:-.02em;margin:0 0 1.2rem;max-width:30ch}
.headline .hl{background:linear-gradient(transparent 58%,var(--mark) 58%,var(--mark) 92%,transparent 92%);padding:0 .1em}
.fields{display:grid;grid-template-columns:1fr 1fr;gap:1rem}
@media (max-width:760px){.fields{grid-template-columns:1fr}}
.field label{display:block;font-weight:600;font-size:.9rem;margin-bottom:.15rem}
.field .hint{color:var(--muted);font-size:.8rem;margin:0 0 .5rem}
.tagbox{min-height:112px;display:flex;flex-wrap:wrap;align-content:flex-start;gap:.4rem;padding:.55rem;border:1px solid var(--line);border-radius:10px;background:var(--sunk);cursor:text}
.tagbox:focus-within{border-color:var(--accent);box-shadow:0 0 0 3px var(--accent-soft)}
.tagbox input{flex:1 1 140px;min-width:120px;border:0;background:transparent;padding:.3rem .2rem;outline:none}
.tag{display:inline-flex;align-items:center;gap:.3rem;background:var(--surface);border:1px solid var(--line);border-radius:7px;padding:.2rem .25rem .2rem .55rem;font-size:.85rem;max-width:100%}
.tag span{overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.tag small{color:var(--muted)}
.tag button{border:0;background:transparent;color:var(--muted);cursor:pointer;width:22px;height:22px;border-radius:5px;display:grid;place-items:center;flex:none}
.tag button:hover{background:var(--sunk);color:var(--ink)}
.names .tag{background:var(--mark);border-color:transparent;color:var(--mark-ink)}
.names .tag small,.names .tag button{color:var(--mark-ink);opacity:.7}

.bottom{display:flex;flex-wrap:wrap;align-items:center;justify-content:space-between;gap:1rem;margin-top:1.1rem}
details.adv summary{cursor:pointer;color:var(--muted);font-size:.88rem;font-weight:500;list-style:none;display:inline-flex;gap:.4rem;align-items:center}
details.adv summary::-webkit-details-marker{display:none}
details.adv summary::before{content:"";width:6px;height:6px;border-right:2px solid currentColor;border-bottom:2px solid currentColor;transform:rotate(-45deg);transition:transform .15s}
details.adv[open] summary::before{transform:rotate(45deg)}
.opts{display:flex;flex-wrap:wrap;gap:1rem 1.5rem;margin-top:.9rem;align-items:end}
.opt label{display:block;font-size:.8rem;color:var(--muted);margin-bottom:.25rem}
.opt input[type=number]{width:96px;padding:.45rem .55rem;border:1px solid var(--line);border-radius:8px;background:var(--sunk)}
.switch{display:flex;align-items:center;gap:.5rem;font-size:.88rem;cursor:pointer}
.switch input{width:1.1rem;height:1.1rem;accent-color:var(--accent)}
.btn{border:1px solid transparent;border-radius:10px;padding:.7rem 1.3rem;font-weight:600;cursor:pointer;display:inline-flex;align-items:center;gap:.5rem}
.btn.primary{background:var(--accent);color:var(--accent-ink)}
.btn.primary:hover{filter:brightness(1.07)}
.btn.quiet{background:var(--surface);border-color:var(--line);padding:.5rem .9rem;font-size:.88rem}
.btn.quiet:hover{background:var(--sunk)}
.btn:disabled{opacity:.45;cursor:not-allowed}
.field.docs{grid-column:1/-1}
.drop{display:flex;flex-wrap:wrap;align-items:center;gap:.4rem;min-height:64px;padding:.55rem;border:1.5px dashed var(--line);border-radius:10px;background:var(--sunk);cursor:pointer}
.drop:hover,.drop.over{border-color:var(--accent);background:var(--accent-soft)}
.drop:focus-within{border-color:var(--accent);box-shadow:0 0 0 3px var(--accent-soft)}
.drop .cta{color:var(--muted);font-size:.88rem;padding:.2rem .3rem}
.drop .cta b{color:var(--accent);font-weight:600}
.sr{position:absolute;width:1px;height:1px;overflow:hidden;clip:rect(0 0 0 0);white-space:nowrap}
.dom.file{color:var(--muted);text-decoration:none}
.formerr{color:var(--warn);font-size:.88rem;margin:.75rem 0 0}

/* progress */
.run{margin-top:1.25rem;display:grid;grid-template-columns:repeat(4,1fr) auto;gap:.75rem;align-items:stretch}
@media (max-width:760px){.run{grid-template-columns:repeat(2,1fr)}}
.stat{background:var(--surface);border:1px solid var(--line);border-radius:12px;padding:.8rem 1rem}
.stat b{display:block;font-size:1.45rem;font-weight:750;font-variant-numeric:tabular-nums;letter-spacing:-.01em}
.stat span{font-size:.8rem;color:var(--muted)}
.stat.flag b{color:var(--accent)}
.runbtns{display:flex;flex-direction:column;gap:.5rem;justify-content:center}
@media (max-width:760px){.runbtns{grid-column:1/-1;flex-direction:row}}
.track{grid-column:1/-1;height:5px;border-radius:3px;background:var(--line);overflow:hidden}
.track div{height:100%;width:0;background:var(--accent);transition:width .4s}
.feed{grid-column:1/-1;font-size:.82rem;color:var(--muted);display:flex;flex-direction:column;gap:.15rem}
.feed div{white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.feed .new{color:var(--ink);font-weight:600}
.feed .new::before{content:"🚩 "}
.errs{grid-column:1/-1;font-size:.82rem}
.errs summary{cursor:pointer;color:var(--warn)}
.errs li{word-break:break-all;color:var(--muted)}

/* results */
.results{margin-top:2rem;display:grid;grid-template-columns:260px 1fr;gap:1.5rem;align-items:start}
@media (max-width:860px){.results{grid-template-columns:1fr}}
.side{position:sticky;top:1rem}
@media (max-width:860px){.side{position:static}}
.side h2,.main h2{font-size:1rem;margin:0 0 .6rem}
.who{display:flex;flex-direction:column;gap:.35rem}
.who button{display:grid;grid-template-columns:1fr auto;gap:.2rem .6rem;text-align:left;padding:.6rem .7rem;border:1px solid transparent;border-radius:10px;background:transparent;cursor:pointer}
.who button:hover{background:var(--surface)}
.who button[aria-pressed=true]{background:var(--surface);border-color:var(--accent);box-shadow:var(--shadow)}
.who .n{font-weight:600;font-size:.9rem;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.who .c{font-variant-numeric:tabular-nums;font-size:.85rem;color:var(--muted)}
.who .b{grid-column:1/-1;height:6px;border-radius:3px;background:var(--sunk);overflow:hidden}
.who .b i{display:block;height:100%;background:var(--mark);border-radius:3px}
.who .zero .n{color:var(--muted);font-weight:500}
.tools{display:flex;flex-wrap:wrap;gap:.6rem;align-items:center;margin-bottom:1rem}
.tools input,.tools select{padding:.5rem .7rem;border:1px solid var(--line);border-radius:9px;background:var(--surface)}
.tools input{flex:1 1 220px}
.card{background:var(--surface);border:1px solid var(--line);border-radius:12px;padding:1rem 1.1rem;margin-bottom:.75rem}
.card.fresh{animation:in .5s ease-out}
@keyframes in{from{background:var(--accent-soft)}}
.card .top{display:flex;justify-content:space-between;gap:1rem;align-items:flex-start}
.card .who-tag{display:inline-block;background:var(--mark);color:var(--mark-ink);font-weight:600;font-size:.8rem;padding:.1rem .5rem;border-radius:5px}
.card .count{font-size:.8rem;color:var(--muted);white-space:nowrap;font-variant-numeric:tabular-nums}
.card h3{font-size:1rem;margin:.45rem 0 .1rem;font-weight:600;overflow-wrap:anywhere}
.card .dom{font-size:.82rem;color:var(--accent);text-decoration:none;overflow-wrap:anywhere}
.card .dom:hover{text-decoration:underline}
.snip{margin:.55rem 0 0;padding-left:.8rem;border-left:2px solid var(--line);font-size:.9rem;color:var(--muted);max-width:78ch;overflow-wrap:anywhere}
.snip mark{background:var(--mark);color:var(--mark-ink);padding:0 .15em;border-radius:3px;font-weight:600}
.meta{display:flex;flex-wrap:wrap;gap:.35rem;margin-top:.7rem}
.meta span{font-size:.75rem;color:var(--muted);background:var(--sunk);border-radius:5px;padding:.1rem .45rem}
.empty{text-align:center;color:var(--muted);padding:3rem 1rem;border:1.5px dashed var(--line);border-radius:14px}
.empty b{display:block;color:var(--ink);font-size:1.05rem;margin-bottom:.25rem}
.hidden{display:none!important}
@media (prefers-reduced-motion:reduce){*{animation:none!important;transition:none!important}}
</style>
</head>
<body>
<div class="shell">
  <header>
    <div class="brand">
      <div class="logo" aria-hidden="true"><svg viewBox="0 0 24 24" fill="none"><path d="M5 21V4" stroke="#fff" stroke-width="2.2" stroke-linecap="round"/><path d="M5 4h11l-2.5 4L16 12H5" fill="#FFE45C"/></svg></div>
      <div><h1>Name Flagger</h1><p>Find where names show up across the web and in your documents</p></div>
    </div>
    <span class="pill" id="pill"><span class="dot"></span><span id="pillText">Ready</span></span>
  </header>

  <section class="setup" aria-label="Crawl setup">
    <p class="headline">Crawl sites or scan your documents, and <span class="hl">flag every name</span> on your list.</p>
    <div class="fields">
      <div class="field">
        <label for="urlIn">Where to start</label>
        <p class="hint">Paste one or more URLs. Press Enter after each. Optional if you add documents.</p>
        <div class="tagbox" id="urlBox"><input id="urlIn" placeholder="example.com/news" autocomplete="off"></div>
      </div>
      <div class="field names">
        <label for="nameIn">Who to look for</label>
        <p class="hint">Press Enter after each name. Add nicknames with |, like Robert Smith | Bob Smith.</p>
        <div class="tagbox" id="nameBox"><input id="nameIn" placeholder="Jane Doe" autocomplete="off"></div>
      </div>
      <div class="field docs">
        <label for="fileIn">Documents to scan</label>
        <p class="hint">PDF, Word (.docx), or text files. Up to {{max_files}} files, {{max_mb}} MB total. Optional if you add URLs.</p>
        <label class="drop" id="drop">
          <input id="fileIn" class="sr" type="file" multiple accept=".pdf,.docx,.txt,.md,.csv,.html,.htm">
          <span class="cta" id="dropCta"><b>Choose files</b> or drag them here</span>
        </label>
      </div>
    </div>
    <div class="bottom">
      <details class="adv">
        <summary>Crawl settings</summary>
        <div class="opts">
          <div class="opt"><label for="depth">Links deep</label><input id="depth" type="number" min="0" max="6" value="2"></div>
          <div class="opt"><label for="maxp">Page limit</label><input id="maxp" type="number" min="1" max="1000" value="200"></div>
          <div class="opt"><label for="delay">Pause per site (s)</label><input id="delay" type="number" min="0" max="10" step="0.1" value="0.5"></div>
          <label class="switch"><input id="ext" type="checkbox" checked> Follow links to other sites</label>
          <label class="switch"><input id="robots" type="checkbox" checked> Respect robots.txt</label>
        </div>
      </details>
      <button class="btn primary" id="go">Start</button>
    </div>
    <p class="formerr" id="formErr" role="alert"></p>
  </section>

  <section class="run hidden" id="run" aria-live="polite">
    <div class="stat"><b id="sPages">0</b><span id="sPagesLbl">pages read</span></div>
    <div class="stat flag"><b id="sFlags">0</b><span>places flagged</span></div>
    <div class="stat"><b id="sQueue">0</b><span>links waiting</span></div>
    <div class="stat"><b id="sTime">0s</b><span>elapsed</span></div>
    <div class="runbtns">
      <button class="btn quiet" id="stop">Stop</button>
      <button class="btn quiet" id="csv">Download CSV</button>
    </div>
    <div class="track"><div id="bar"></div></div>
    <div class="feed" id="feed"></div>
    <details class="errs hidden" id="errs"><summary id="errSum"></summary><ul id="errList"></ul></details>
  </section>

  <section class="results" id="results">
    <aside class="side">
      <h2>Names</h2>
      <div class="who" id="who"><p style="color:var(--muted);font-size:.88rem;margin:0">Names you add appear here with how many pages or files mention them.</p></div>
    </aside>
    <div class="main">
      <div class="tools hidden" id="tools">
        <input id="q" type="search" placeholder="Filter by page, site, or text" aria-label="Filter results">
        <select id="sort" aria-label="Sort results">
          <option value="new">Newest first</option>
          <option value="most">Most mentions</option>
          <option value="name">By name</option>
          <option value="site">By site</option>
        </select>
      </div>
      <div id="list"><div class="empty"><b>No results yet</b>Add a starting URL or some documents, plus at least one name, then press Start. Flagged pages and files show up here as they're found.</div></div>
    </div>
  </section>
</div>

<script>
const $ = id => document.getElementById(id);
const esc = s => String(s).replace(/[&<>"']/g, c => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));
const host = u => { try { return new URL(u).hostname.replace(/^www\./, ''); } catch { return u; } };
const X = '<svg width="12" height="12" viewBox="0 0 12 12" aria-hidden="true"><path d="M3 3l6 6M9 3l-6 6" stroke="currentColor" stroke-width="1.6" stroke-linecap="round"/></svg>';

/* ---------- tag inputs ---------- */
function tagInput(box, input, {split, label}) {
  const items = [];
  const ph = input.placeholder;
  const draw = () => {
    input.placeholder = items.length ? 'Add another…' : ph;
    box.querySelectorAll('.tag').forEach(t => t.remove());
    items.forEach((v, i) => {
      const t = document.createElement('span'); t.className = 'tag';
      t.innerHTML = label(v) + `<button type="button" aria-label="Remove ${esc(v)}">${X}</button>`;
      t.querySelector('button').onclick = () => { items.splice(i, 1); draw(); input.focus(); };
      box.insertBefore(t, input);
    });
  };
  const add = raw => { raw.split(split).map(s => s.trim()).filter(Boolean).forEach(v => { if (!items.includes(v)) items.push(v); }); draw(); };
  input.addEventListener('keydown', e => {
    if ((e.key === 'Enter' || e.key === ',') && input.value.trim()) { e.preventDefault(); add(input.value); input.value = ''; }
    else if (e.key === 'Backspace' && !input.value && items.length) { items.pop(); draw(); }
  });
  input.addEventListener('paste', e => { const t = e.clipboardData.getData('text'); if (/[\n,]/.test(t)) { e.preventDefault(); add(t); } });
  input.addEventListener('blur', () => { if (input.value.trim()) { add(input.value); input.value = ''; } });
  box.addEventListener('click', e => { if (e.target === box) input.focus(); });
  return { get: () => items.slice(), commit: () => { if (input.value.trim()) { add(input.value); input.value = ''; } } };
}
const urls = tagInput($('urlBox'), $('urlIn'), { split: /[\s,]+/, label: v => `<span>${esc(v)}</span>` });
const names = tagInput($('nameBox'), $('nameIn'), { split: /[\n,]+/, label: v => {
  const [main, ...alt] = v.split('|').map(s => s.trim()).filter(Boolean);
  return `<span>${esc(main)}${alt.length ? ` <small>also ${esc(alt.join(', '))}</small>` : ''}</span>`;
}});

/* ---------- file picker ---------- */
let docs = [];
const MAX_FILES = {{max_files}}, MAX_BYTES = {{max_mb}} * 1024 * 1024;
const kb = n => n < 1048576 ? Math.max(1, Math.round(n / 1024)) + ' KB' : (n / 1048576).toFixed(1) + ' MB';
function drawDocs() {
  const drop = $('drop');
  drop.querySelectorAll('.tag').forEach(t => t.remove());
  docs.forEach((f, i) => {
    const t = document.createElement('span'); t.className = 'tag';
    t.innerHTML = `<span>${esc(f.name)} <small>${kb(f.size)}</small></span><button type="button" aria-label="Remove ${esc(f.name)}">${X}</button>`;
    t.querySelector('button').onclick = e => { e.preventDefault(); e.stopPropagation(); docs.splice(i, 1); drawDocs(); };
    drop.insertBefore(t, $('dropCta'));
  });
  $('dropCta').innerHTML = docs.length ? '<b>Add more</b>' : '<b>Choose files</b> or drag them here';
}
function addDocs(list) {
  for (const f of list) if (!docs.some(d => d.name === f.name && d.size === f.size)) docs.push(f);
  drawDocs();
}
$('fileIn').onchange = e => { addDocs(e.target.files); e.target.value = ''; };
['dragenter', 'dragover'].forEach(ev => $('drop').addEventListener(ev, e => { e.preventDefault(); $('drop').classList.add('over'); }));
['dragleave', 'drop'].forEach(ev => $('drop').addEventListener(ev, e => { e.preventDefault(); $('drop').classList.remove('over'); }));
$('drop').addEventListener('drop', e => addDocs(e.dataTransfer.files));

/* ---------- crawl control ---------- */
let jobId = null, timer = null, last = null, filter = null, seenHits = new Set();

$('go').onclick = async () => {
  urls.commit(); names.commit();
  $('formErr').textContent = '';
  if (!urls.get().length && !docs.length) { $('formErr').textContent = 'Add a website to start from or a document to scan.'; $('urlIn').focus(); return; }
  if (docs.length > MAX_FILES) { $('formErr').textContent = `You can scan up to ${MAX_FILES} files at a time.`; return; }
  if (docs.reduce((a, f) => a + f.size, 0) > MAX_BYTES) { $('formErr').textContent = `Those files are too big together. The limit is ${MAX_BYTES / 1048576} MB.`; return; }
  if (!names.get().length) { $('formErr').textContent = 'Add at least one name to look for.'; $('nameIn').focus(); return; }
  $('go').disabled = true;
  try {
    const fd = new FormData();
    Object.entries({ urls: urls.get().join('\n'), names: names.get().join('\n'),
      max_depth: $('depth').value, max_pages: $('maxp').value, delay: $('delay').value,
      follow_external: $('ext').checked, respect_robots: $('robots').checked }).forEach(([k, v]) => fd.append(k, v));
    docs.forEach(f => fd.append('files', f, f.name));
    const r = await fetch('start', { method: 'POST', body: fd });
    let d; try { d = await r.json(); } catch { d = { error: r.status === 413 ? 'Those files are too big to upload.' : 'The server didn\'t respond properly.' }; }
    if (!r.ok) throw new Error(d.error || 'The crawl couldn\'t start.');
    jobId = d.job_id; filter = null; seenHits = new Set();
    $('run').classList.remove('hidden'); $('stop').disabled = false;
    clearInterval(timer); poll(); timer = setInterval(poll, 1000);
  } catch (e) { $('formErr').textContent = e.message; $('go').disabled = false; }
};
$('stop').onclick = () => { if (jobId) { fetch('stop/' + jobId, { method: 'POST' }); $('stop').disabled = true; } };
$('csv').onclick = () => { if (jobId) location.href = 'export/' + jobId + '.csv'; };
$('q').oninput = () => last && drawList(last);
$('sort').onchange = () => last && drawList(last);

async function poll() {
  try {
    const r = await fetch('status/' + jobId); const d = await r.json();
    if (!r.ok) throw new Error(d.error);
    last = d; draw(d);
    if (d.status !== 'running') { clearInterval(timer); $('go').disabled = false; $('stop').disabled = true; }
  } catch (e) { clearInterval(timer); setPill('error', 'Lost connection'); $('go').disabled = false; }
}

function setPill(cls, text) { $('pill').className = 'pill ' + cls; $('pillText').textContent = text; }

/* ---------- rendering ---------- */
function draw(d) {
  const labels = { running: d.files_scanned < d.files_total ? 'Scanning files' : 'Crawling', done: 'Finished', stopped: 'Stopped', error: 'Stopped with an error' };
  setPill(d.status === 'running' ? 'running' : d.status === 'error' ? 'error' : 'done', labels[d.status] || d.status);
  $('sPages').textContent = d.pages_crawled + d.files_scanned;
  $('sPagesLbl').textContent = !d.files_total ? 'pages read' : !d.start_urls.length ? 'files read' : 'pages & files read';
  $('sFlags').textContent = d.hits.length;
  $('sQueue').textContent = d.queued;
  $('sTime').textContent = d.elapsed < 60 ? Math.round(d.elapsed) + 's' : Math.floor(d.elapsed / 60) + 'm ' + Math.round(d.elapsed % 60) + 's';
  $('bar').style.width = (d.status === 'running' ? Math.min(100, 100 * (d.pages_crawled + d.files_scanned) / Math.max(1, d.progress_total)) : 100) + '%';
  const path = u => { try { return new URL(u).pathname; } catch { return ''; } };
  const feed = (d.current ? [`<div>Reading ${esc(d.current)}</div>`] : [])
    .concat(d.recent.slice(0, 3).map(p => `<div class="${p.new_flags ? 'new' : ''}">${p.new_flags ? `Flagged ${esc(p.title || host(p.url))}` : p.file ? `Scanned ${esc(p.url)}` : `Read ${esc(host(p.url))}${esc(path(p.url))}`}</div>`));
  $('feed').innerHTML = feed.join('');
  $('errs').classList.toggle('hidden', !d.error_count);
  $('errSum').textContent = `${d.error_count} page${d.error_count === 1 ? '' : 's'} or file${d.error_count === 1 ? '' : 's'} couldn't be read` + (d.skipped_robots ? `, ${d.skipped_robots} blocked by robots.txt` : '');
  $('errList').innerHTML = d.errors.map(e => `<li>${esc(e.url)}: ${esc(e.error)}</li>`).join('');
  drawWho(d); drawList(d);
}

function drawWho(d) {
  const pages = {}, mentions = {};
  d.names.forEach(n => { pages[n] = 0; mentions[n] = 0; });
  d.hits.forEach(h => { pages[h.name]++; mentions[h.name] += h.count; });
  const max = Math.max(1, ...Object.values(pages));
  const unit = d.start_urls.length ? 'place' : 'file';
  const btn = (key, label, p, m) => `<button data-n="${esc(key)}" aria-pressed="${filter === key}" class="${p ? '' : 'zero'}">
      <span class="n">${esc(label)}</span><span class="c">${p} ${unit}${p === 1 ? '' : 's'}</span>
      ${key === '' ? '' : `<span class="b"><i style="width:${100 * p / max}%"></i></span>`}</button>`;
  $('who').innerHTML = btn('', 'All names', d.hits.length, 0).replace('aria-pressed="false"', `aria-pressed="${filter === null}"`)
    + [...d.names].sort((a, b) => pages[b] - pages[a]).map(n => btn(n, n, pages[n], mentions[n])).join('');
  $('who').querySelectorAll('button').forEach(b => b.onclick = () => { filter = b.dataset.n || null; drawWho(d); drawList(d); });
}

function drawList(d) {
  $('tools').classList.toggle('hidden', !d.hits.length);
  const q = $('q').value.trim().toLowerCase();
  let hits = d.hits.map((h, i) => ({ ...h, i })).filter(h => !filter || h.name === filter);
  if (q) hits = hits.filter(h => (h.title + ' ' + h.url + ' ' + h.snippets.map(s => s.pre + s.match + s.post).join(' ')).toLowerCase().includes(q));
  const s = $('sort').value;
  hits.sort(s === 'most' ? (a, b) => b.count - a.count : s === 'name' ? (a, b) => a.name.localeCompare(b.name) || b.count - a.count
    : s === 'site' ? (a, b) => host(a.url).localeCompare(host(b.url)) : (a, b) => b.i - a.i);

  if (!hits.length) {
    const running = d.status === 'running', onlyFiles = !d.start_urls.length;
    $('list').innerHTML = `<div class="empty"><b>${d.hits.length ? 'Nothing matches that filter' : running ? 'Looking for names…' : 'No mentions found'}</b>${
      d.hits.length ? 'Try a different name or clear the search box.' : running ? 'Flagged pages and files will appear here as soon as they\'re found.' : onlyFiles ? 'Try adding nicknames or other spellings. Scanned-image PDFs have no text to search.' : 'Try going more links deep, raising the page limit, or adding nicknames.'}</div>`;
    return;
  }
  $('list').innerHTML = hits.map(h => {
    const key = h.name + '|' + h.url, fresh = !seenHits.has(key); seenHits.add(key);
    return `<article class="card ${fresh ? 'fresh' : ''}">
      <div class="top"><span class="who-tag">${esc(h.name)}</span><span class="count">${h.count} mention${h.count === 1 ? '' : 's'}</span></div>
      <h3>${esc(h.title || host(h.url))}</h3>
      ${h.file ? `<span class="dom file">Uploaded file</span>` : `<a class="dom" href="${esc(h.url)}" target="_blank" rel="noopener">${esc(h.url)}</a>`}
      ${h.snippets.map(sn => `<p class="snip">${h.file && sn.where !== 'document' ? `<small>${esc(sn.where)}:</small> ` : ''}${esc(sn.pre)}<mark>${esc(sn.match)}</mark>${esc(sn.post)}</p>`).join('')}
      <div class="meta">${h.variants.map(v => `<span>${esc(v)}</span>`).join('')}${h.where.includes('link URL') ? '<span>in a link</span>' : ''}${h.file && h.where[0] !== 'document' ? `<span>${esc(h.where.slice(0, 6).join(', '))}${h.where.length > 6 ? '…' : ''}</span>` : ''}</div>
    </article>`;
  }).join('');
}
</script>
</body>
</html>
"""

if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5000"))
    app.run(host="0.0.0.0" if "PORT" in os.environ else "127.0.0.1",
            port=port, debug=False, threaded=True)
